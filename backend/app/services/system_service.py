import os
import io
import csv
import json
import zipfile
import shutil
from datetime import datetime, timezone, date
from decimal import Decimal
from typing import Dict, Any, List, Optional, Tuple
from sqlalchemy import text, select, delete
from sqlalchemy.orm import Session
from fastapi import HTTPException, status, UploadFile

from app.core.config import settings
from app.models.setting import SystemSetting
from app.models.product import Product, ProductImage, ProductKitItem
from app.models.category import Category
from app.models.unit import UnitOfMeasure
from app.models.color import ProductColor
from app.models.material import ProductMaterial
from app.models.supply import ProductSupply
from app.models.item_type import ItemTypeModel
from app.models.stock import StockMovement, MovementType
from app.models.sale import Customer, Sale, SaleItem, SaleStatus
from app.models.production import (
    ProductionRecipe,
    ProductionRecipeItem,
    ProductionOrder,
    ProductionOrderItem,
)
from app.models.financial import FinancialCategory, FinancialTransaction

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class SystemService:
    @staticmethod
    def seed_default_settings(db: Session) -> None:
        """Inicializa as configurações padrão do sistema caso ainda não existam."""
        defaults = {
            "company_info": json.dumps({
                "trade_name": "Ateliê Casa Luma",
                "legal_name": "Ateliê Casa Luma Decorações Artesanais",
                "document": "",
                "state_registration": "",
                "phone": "(31) 98337-1642",
                "email": "contato@casaluma.com.br",
                "address_line": "",
                "city": "Belo Horizonte",
                "state": "MG",
                "zip_code": "",
                "pix_key": "31983371642",
                "pix_key_type": "Celular",
                "pix_bank": "Nubank",
                "pix_owner_name": "Ateliê Casa Luma",
                "logo_url": None,
            }, ensure_ascii=False),
            "whatsapp_templates": json.dumps({
                "order_confirmation": "Olá, {cliente}! 🌿\nSeu pedido #{pedido} no Ateliê Casa Luma foi confirmado no valor de {valor}.\nJá estamos preparando tudo com muito carinho!",
                "pix_billing": "Olá, {cliente}! ✨\nSegue a chave Pix para pagamento do seu pedido #{pedido}:\n🔑 Chave Pix ({tipo}): {chave_pix}\n🏦 Banco: {banco}\n👤 Titular: {titular}\n💰 Valor total: {valor}",
                "shipping_notification": "Olá, {cliente}! 📦✨\nSuas peças artesanais do Ateliê Casa Luma já foram enviadas!\nCódigo de rastreio: {rastreio}\nObrigada por apoiar o artesanato autoral!",
                "follow_up": "Olá, {cliente}! 🌸\nTudo bem? Suas peças chegaram com segurança? Adoraríamos ver como ficaram no seu ambiente! Se puder, marque nosso Instagram @ateliecasaluma 📸",
            }, ensure_ascii=False),
            "studio_policies": json.dumps({
                "production_cure_time": "3 a 5 dias úteis para cura natural do cimento e maturação das essências",
                "care_instructions": "Peças em cimento: limpar com pano seco ou levemente umedecido; não utilizar abrasivos. Velas: aparar o pavio a 5mm antes de cada queima e deixar a cera derreter até a borda na primeira queima.",
                "warranty_terms": "Garantia artesanal de 30 dias para eventuais defeitos de fabricação.",
            }, ensure_ascii=False),
        }

        for k, v in defaults.items():
            existing = db.get(SystemSetting, k)
            if not existing:
                setting = SystemSetting(key=k, value=v, description=f"Configurações de {k}")
                db.add(setting)
        db.commit()

    @staticmethod
    def get_all_settings(db: Session) -> Dict[str, Any]:
        """Retorna todas as configurações do sistema em formato dicionário."""
        settings_rows = db.query(SystemSetting).all()
        result: Dict[str, Any] = {}
        for row in settings_rows:
            try:
                result[row.key] = json.loads(row.value)
            except Exception:
                result[row.key] = row.value
        return result

    @staticmethod
    def update_setting(db: Session, key: str, value: Any, description: Optional[str] = None) -> SystemSetting:
        """Salva ou atualiza uma configuração do sistema."""
        setting = db.get(SystemSetting, key)
        json_val = json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value

        if setting:
            setting.value = json_val
            if description:
                setting.description = description
            setting.updated_at = datetime.now(timezone.utc)
        else:
            setting = SystemSetting(
                key=key,
                value=json_val,
                description=description,
                updated_at=datetime.now(timezone.utc),
            )
            db.add(setting)
        db.commit()
        db.refresh(setting)
        return setting

    # --- Backup e Restore ---
    @staticmethod
    def export_backup_zip(db: Session) -> io.BytesIO:
        """Gera um arquivo .zip contendo database_dump.json, database_dump.sql e a pasta de mídia."""
        # 1. Serializar dados de todas as tabelas
        data = {
            "version": "1.0",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "tables": {
                "categories": [SystemService._row_to_dict(r) for r in db.query(Category).all()],
                "units_of_measure": [SystemService._row_to_dict(r) for r in db.query(UnitOfMeasure).all()],
                "item_types": [SystemService._row_to_dict(r) for r in db.query(ItemTypeModel).all()],
                "product_colors": [SystemService._row_to_dict(r) for r in db.query(ProductColor).all()],
                "product_materials": [SystemService._row_to_dict(r) for r in db.query(ProductMaterial).all()],
                "product_supplies": [SystemService._row_to_dict(r) for r in db.query(ProductSupply).all()],
                "products": [SystemService._row_to_dict(r) for r in db.query(Product).all()],
                "product_images": [SystemService._row_to_dict(r) for r in db.query(ProductImage).all()],
                "product_kit_items": [SystemService._row_to_dict(r) for r in db.query(ProductKitItem).all()],
                "stock_movements": [SystemService._row_to_dict(r) for r in db.query(StockMovement).all()],
                "customers": [SystemService._row_to_dict(r) for r in db.query(Customer).all()],
                "sales": [SystemService._row_to_dict(r) for r in db.query(Sale).all()],
                "sale_items": [SystemService._row_to_dict(r) for r in db.query(SaleItem).all()],
                "production_recipes": [SystemService._row_to_dict(r) for r in db.query(ProductionRecipe).all()],
                "production_recipe_items": [SystemService._row_to_dict(r) for r in db.query(ProductionRecipeItem).all()],
                "production_orders": [SystemService._row_to_dict(r) for r in db.query(ProductionOrder).all()],
                "production_order_items": [SystemService._row_to_dict(r) for r in db.query(ProductionOrderItem).all()],
                "financial_categories": [SystemService._row_to_dict(r) for r in db.query(FinancialCategory).all()],
                "financial_transactions": [SystemService._row_to_dict(r) for r in db.query(FinancialTransaction).all()],
                "system_settings": [SystemService._row_to_dict(r) for r in db.query(SystemSetting).all()],
            }
        }

        # 2. Gerar arquivo SQL equivalente
        sql_lines = ["-- Backup do Ateliê Casa Luma Estoque", f"-- Gerado em: {datetime.now(timezone.utc).isoformat()}", "SET FOREIGN_KEY_CHECKS=0;\n"]
        for table_name, rows in data["tables"].items():
            if rows:
                cols = list(rows[0].keys())
                cols_str = ", ".join([f"`{c}`" for c in cols])
                for r in rows:
                    vals = []
                    for c in cols:
                        v = r[c]
                        if v is None:
                            vals.append("NULL")
                        elif isinstance(v, (int, float, Decimal)):
                            vals.append(str(v))
                        elif isinstance(v, bool):
                            vals.append("1" if v else "0")
                        else:
                            clean_v = str(v).replace("\\", "\\\\").replace("'", "''")
                            vals.append(f"'{clean_v}'")
                    sql_lines.append(f"INSERT INTO `{table_name}` ({cols_str}) VALUES ({', '.join(vals)});")
        sql_lines.append("\nSET FOREIGN_KEY_CHECKS=1;\n")
        sql_content = "\n".join(sql_lines)

        # 3. Compactar em ZIP
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            zip_file.writestr("database_dump.json", json.dumps(data, indent=2, ensure_ascii=False))
            zip_file.writestr("database_dump.sql", sql_content)
            
            # Adicionar pasta media se existir
            media_dir = getattr(settings, "MEDIA_DIR", "/app/media")
            if os.path.exists(media_dir):
                for root, _, files in os.walk(media_dir):
                    for file in files:
                        full_path = os.path.join(root, file)
                        rel_path = os.path.relpath(full_path, media_dir)
                        zip_file.write(full_path, arcname=f"media/{rel_path}")

        zip_buffer.seek(0)
        return zip_buffer

    @staticmethod
    def restore_backup(db: Session, file_bytes: bytes, filename: str, confirmation_code: str) -> Dict[str, Any]:
        """Restaura o banco de dados a partir de um arquivo .zip ou .json com confirmação obrigatória."""
        if confirmation_code.strip() != "RESTAURAR":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Código de confirmação incorreto. Digite 'RESTAURAR' para prosseguir.",
            )

        data = None
        media_files_to_extract = []

        # Detectar se é ZIP ou JSON
        if filename.endswith(".zip") or file_bytes.startswith(b"PK\x03\x04"):
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as zf:
                    if "database_dump.json" in zf.namelist():
                        data = json.loads(zf.read("database_dump.json").decode("utf-8"))
                    for name in zf.namelist():
                        if name.startswith("media/") and not name.endswith("/"):
                            media_files_to_extract.append((name[6:], zf.read(name)))
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Arquivo ZIP corrompido ou inválido: {e}")
        elif filename.endswith(".json") or file_bytes.strip().startswith(b"{"):
            try:
                data = json.loads(file_bytes.decode("utf-8"))
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Arquivo JSON corrompido ou inválido: {e}")
        else:
            raise HTTPException(status_code=400, detail="Formato de arquivo não suportado. Envie um arquivo .zip ou .json.")

        if not data or "tables" not in data:
            raise HTTPException(status_code=400, detail="Arquivo de backup inválido: dados de tabelas não encontrados.")

        tables = data["tables"]
        restored_counts = {}

        # Execução atômica no banco de dados
        try:
            db.execute(text("SET FOREIGN_KEY_CHECKS = 0;"))

            # Limpar tabelas existentes
            order_to_clear = [
                "production_order_items", "production_orders", "production_recipe_items", "production_recipes",
                "sale_items", "sales", "stock_movements", "product_kit_items", "product_images",
                "products", "product_supplies", "product_materials", "product_colors", "item_types",
                "units_of_measure", "categories", "customers", "financial_transactions", "financial_categories", "system_settings"
            ]
            for tbl in order_to_clear:
                db.execute(text(f"DELETE FROM `{tbl}`;"))

            # Inserir tabelas na ordem de dependência
            order_to_insert = [
                ("categories", Category),
                ("units_of_measure", UnitOfMeasure),
                ("item_types", ItemTypeModel),
                ("product_colors", ProductColor),
                ("product_materials", ProductMaterial),
                ("product_supplies", ProductSupply),
                ("products", Product),
                ("product_images", ProductImage),
                ("product_kit_items", ProductKitItem),
                ("stock_movements", StockMovement),
                ("customers", Customer),
                ("sales", Sale),
                ("sale_items", SaleItem),
                ("production_recipes", ProductionRecipe),
                ("production_recipe_items", ProductionRecipeItem),
                ("production_orders", ProductionOrder),
                ("production_order_items", ProductionOrderItem),
                ("financial_categories", FinancialCategory),
                ("financial_transactions", FinancialTransaction),
                ("system_settings", SystemSetting),
            ]

            for tbl_name, model_class in order_to_insert:
                rows = tables.get(tbl_name, [])
                restored_counts[tbl_name] = len(rows)
                for r in rows:
                    cols = list(r.keys())
                    cols_str = ", ".join([f"`{c}`" for c in cols])
                    params_str = ", ".join([f":{c}" for c in cols])
                    db.execute(text(f"INSERT INTO `{tbl_name}` ({cols_str}) VALUES ({params_str});"), r)

            db.execute(text("SET FOREIGN_KEY_CHECKS = 1;"))
            db.commit()

            # Extrair mídia se houver
            media_dir = getattr(settings, "MEDIA_DIR", "/app/media")
            if media_files_to_extract and os.path.exists(media_dir):
                for rel_path, content in media_files_to_extract:
                    dest_path = os.path.join(media_dir, rel_path)
                    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
                    with open(dest_path, "wb") as f:
                        f.write(content)

            return {
                "success": True,
                "message": "Backup restaurado com sucesso!",
                "restored_records": restored_counts,
                "extracted_media_files": len(media_files_to_extract),
            }
        except Exception as err:
            db.rollback()
            try:
                db.execute(text("SET FOREIGN_KEY_CHECKS = 1;"))
            except Exception:
                pass
            raise HTTPException(status_code=500, detail=f"Erro ao restaurar banco de dados: {err}")

    # --- Exportação em Excel (.xlsx) e CSV (.csv) ---
    @staticmethod
    def export_products(db: Session, export_format: str = "xlsx") -> Tuple[io.BytesIO, str, str]:
        """Exporta o catálogo de produtos e variações em Excel ou CSV."""
        from app.services.stock_service import stock_service

        products = db.query(Product).order_by(Product.id.asc()).all()
        headers = [
            "ID", "SKU", "Nome do Produto", "Tipo de Item", "Categoria", "Cor", "Tamanho / Volume",
            "Material", "Preço de Venda (R$)", "Custo Unitário (R$)", "Estoque Atual", "Estoque Mínimo",
            "Ativo", "É Variação", "SKU do Pai", "Nome da Variação"
        ]

        rows = []
        for p in products:
            bal = stock_service.get_product_balance(db, p.id)
            rows.append([
                p.id,
                p.sku,
                p.name,
                p.item_type,
                p.category.name if p.category else "",
                p.color or "",
                p.size or "",
                p.material or "",
                float(p.sale_price) if p.sale_price is not None else 0.0,
                float(p.unit_cost) if p.unit_cost is not None else 0.0,
                float(bal),
                float(p.min_stock),
                "Sim" if p.is_active else "Não",
                "Sim" if p.parent_id else "Não",
                p.parent.sku if p.parent else "",
                p.variation_name or "",
            ])

        return SystemService._create_spreadsheet(
            title="Produtos e Variações",
            headers=headers,
            rows=rows,
            export_format=export_format,
            filename_prefix="produtos_atelie_casaluma"
        )

    @staticmethod
    def export_customers(db: Session, export_format: str = "xlsx") -> Tuple[io.BytesIO, str, str]:
        """Exporta os clientes do CRM em Excel ou CSV."""
        customers = db.query(Customer).order_by(Customer.name.asc()).all()
        headers = [
            "ID", "Nome Completo", "Apelido", "Telefone / WhatsApp", "E-mail",
            "Cidade", "Estado", "Endereço", "Bairro", "CEP", "Canal de Origem",
            "Data de Nascimento", "Tags", "Cadastrado em"
        ]

        rows = []
        for c in customers:
            rows.append([
                c.id,
                c.name,
                getattr(c, "nickname", "") or "",
                c.phone or "",
                c.email or "",
                getattr(c, "city", "") or "",
                getattr(c, "state", "") or "",
                getattr(c, "address_line", "") or "",
                getattr(c, "neighborhood", "") or "",
                getattr(c, "zip_code", "") or "",
                getattr(c, "origin_channel", "") or "",
                getattr(c, "birth_date", "") or "",
                getattr(c, "tags", "") or "",
                c.created_at.strftime("%d/%m/%Y %H:%M") if hasattr(c, "created_at") and c.created_at else "",
            ])

        return SystemService._create_spreadsheet(
            title="Clientes (CRM)",
            headers=headers,
            rows=rows,
            export_format=export_format,
            filename_prefix="clientes_crm_casaluma"
        )

    @staticmethod
    def export_sales(db: Session, export_format: str = "xlsx") -> Tuple[io.BytesIO, str, str]:
        """Exporta o histórico de vendas com detalhes em Excel ou CSV."""
        sales = db.query(Sale).order_by(Sale.created_at.desc()).all()
        headers = [
            "ID Venda", "Data / Hora", "Cliente", "Telefone", "Status",
            "Forma de Pagamento", "Valor Total (R$)", "Quantidade de Itens", "Itens do Pedido"
        ]

        rows = []
        for s in sales:
            items_desc = "; ".join([f"{it.product.name} (x{it.quantity})" for it in s.items if it.product])
            rows.append([
                s.id,
                s.created_at.strftime("%d/%m/%Y %H:%M") if s.created_at else "",
                s.customer.name if s.customer else "Cliente Avulso",
                s.customer.phone if s.customer and s.customer.phone else "",
                s.status.value if hasattr(s.status, "value") else str(s.status),
                getattr(s, "payment_method", "") or "Não informado",
                float(s.total_amount) if s.total_amount is not None else 0.0,
                len(s.items),
                items_desc,
            ])

        return SystemService._create_spreadsheet(
            title="Histórico de Vendas",
            headers=headers,
            rows=rows,
            export_format=export_format,
            filename_prefix="vendas_atelie_casaluma"
        )

    @staticmethod
    def export_stock(db: Session, export_format: str = "xlsx") -> Tuple[io.BytesIO, str, str]:
        """Exporta o balanço de estoque detalhado em Excel ou CSV."""
        from app.services.stock_service import stock_service

        products = db.query(Product).order_by(Product.name.asc()).all()
        headers = [
            "ID", "SKU", "Produto", "Tipo", "Categoria", "Unidade", "Cor", "Tamanho",
            "Estoque Atual", "Estoque Mínimo", "Custo Unitário (R$)", "Valor Imobilizado (R$)"
        ]

        rows = []
        for p in products:
            bal = stock_service.get_product_balance(db, p.id)
            total_cost = float(bal) * float(p.unit_cost)
            rows.append([
                p.id,
                p.sku,
                p.name + (f" - {p.variation_name}" if p.variation_name else ""),
                p.item_type,
                p.category.name if p.category else "",
                p.unit.code if p.unit else "UN",
                p.color or "",
                p.size or "",
                float(bal),
                float(p.min_stock),
                float(p.unit_cost),
                round(total_cost, 2),
            ])

        return SystemService._create_spreadsheet(
            title="Inventário de Estoque",
            headers=headers,
            rows=rows,
            export_format=export_format,
            filename_prefix="estoque_inventario_casaluma"
        )

    # --- Planilhas Modelo para Importação ---
    @staticmethod
    def get_template(template_type: str, export_format: str = "xlsx") -> Tuple[io.BytesIO, str, str]:
        """Retorna uma planilha modelo formatada para produtos ou clientes."""
        if template_type == "products":
            headers = [
                "SKU", "Nome do Produto", "Tipo de Item", "Categoria", "Cor", "Tamanho / Volume",
                "Material", "Preço de Venda (R$)", "Custo Unitário (R$)", "Estoque Inicial", "Estoque Mínimo"
            ]
            rows = [
                ["VELA-LAV-180G", "Vela Aromática Lavanda Francesa", "PRODUTO_ACABADO", "Velas Aromáticas", "Lilás", "180g", "Cera de Coco", 49.90, 16.50, 15, 3],
                ["BAN-NUV-BRAN", "Bandeja Nuvem Decorativa", "PRODUTO_ACABADO", "Bandejas", "Branco Marmorizado", "P", "Cimento + Argamassa", 39.90, 12.00, 8, 2],
                ["ESS-LAV-1L", "Essência Pura Lavanda", "MATERIA_PRIMA", "Insumos", "", "1 Litro", "", 0.00, 140.00, 2, 1],
            ]
            return SystemService._create_spreadsheet(
                title="Modelo Importação Produtos",
                headers=headers,
                rows=rows,
                export_format=export_format,
                filename_prefix="modelo_importacao_produtos"
            )
        elif template_type == "customers":
            headers = [
                "Nome Completo", "Apelido", "Telefone / WhatsApp", "E-mail", "Cidade", "Estado",
                "Endereço", "Bairro", "CEP", "Canal de Origem", "Data de Nascimento"
            ]
            rows = [
                ["Marina Silva", "Mari", "(31) 98765-4321", "marina@exemplo.com", "Belo Horizonte", "MG", "Rua das Flores, 120", "Savassi", "30110-000", "Instagram", "15/04/1990"],
                ["Lucas Souza", "", "(31) 99887-1122", "lucas@exemplo.com", "Contagem", "MG", "Av. Brasil, 450", "Eldorado", "32310-000", "Feira de Artesanato", "22/11/1988"],
            ]
            return SystemService._create_spreadsheet(
                title="Modelo Importação Clientes",
                headers=headers,
                rows=rows,
                export_format=export_format,
                filename_prefix="modelo_importacao_clientes"
            )
        else:
            raise HTTPException(status_code=400, detail="Tipo de modelo desconhecido.")

    # --- Importação de Dados ---
    @staticmethod
    def import_products(db: Session, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Lê arquivo Excel (.xlsx) ou CSV e realiza a importação de produtos no banco."""
        rows = SystemService._read_spreadsheet(file_content, filename)
        if not rows:
            raise HTTPException(status_code=400, detail="Planilha vazia ou sem dados para importar.")

        from app.services.product_service import product_service
        from app.schemas.product import ProductCreate

        # Buscar ou criar unidades e categorias auxiliares padrão
        default_unit = db.query(UnitOfMeasure).filter(UnitOfMeasure.code == "UN").first()
        if not default_unit:
            default_unit = db.query(UnitOfMeasure).first()

        created_count = 0
        updated_count = 0
        errors = []

        for idx, row in enumerate(rows, start=2):
            try:
                sku = str(row.get("SKU") or "").strip().upper()
                name = str(row.get("Nome do Produto") or "").strip()
                if not sku or not name:
                    continue

                item_type = str(row.get("Tipo de Item") or "PRODUTO_ACABADO").strip().upper()
                cat_name = str(row.get("Categoria") or "").strip()
                color = str(row.get("Cor") or "").strip() or None
                size = str(row.get("Tamanho / Volume") or "").strip() or None
                material = str(row.get("Material") or "").strip() or None

                sale_price = SystemService._parse_decimal(row.get("Preço de Venda (R$)"))
                unit_cost = SystemService._parse_decimal(row.get("Custo Unitário (R$)")) or Decimal("0.00")
                initial_stock = SystemService._parse_decimal(row.get("Estoque Inicial")) or Decimal("0.00")
                min_stock = SystemService._parse_decimal(row.get("Estoque Mínimo")) or Decimal("0.00")

                # Categoria
                cat_id = None
                if cat_name:
                    cat = db.query(Category).filter(Category.name.ilike(cat_name)).first()
                    if not cat:
                        cat = Category(name=cat_name, description="Importada via planilha")
                        db.add(cat)
                        db.flush()
                    cat_id = cat.id

                # Verificar se já existe pelo SKU
                existing = ProductService_get = db.query(Product).filter(Product.sku == sku).first()
                if existing:
                    existing.name = name
                    existing.item_type = item_type
                    if cat_id: existing.category_id = cat_id
                    if color: existing.color = color
                    if size: existing.size = size
                    if material: existing.material = material
                    if sale_price is not None: existing.sale_price = sale_price
                    existing.unit_cost = unit_cost
                    existing.min_stock = min_stock
                    updated_count += 1
                else:
                    p_in = ProductCreate(
                        sku=sku,
                        name=name,
                        item_type=item_type,
                        category_id=cat_id,
                        unit_id=default_unit.id if default_unit else 1,
                        unit_cost=unit_cost,
                        sale_price=sale_price,
                        color=color,
                        size=size,
                        material=material,
                        min_stock=min_stock,
                    )
                    created_prod = product_service.create_product(db, p_in)

                    # Lançar estoque inicial se informado
                    if initial_stock > 0:
                        from app.models.stock import StockMovement, MovementType
                        mov = StockMovement(
                            product_id=created_prod.id,
                            movement_type=MovementType.ENTRADA_INICIAL,
                            quantity=initial_stock,
                            unit_cost_at_moment=unit_cost,
                            justification="Estoque inicial via importação de planilha",
                        )
                        db.add(mov)

                    created_count += 1
            except Exception as e:
                errors.append(f"Linha {idx} ({row.get('SKU', 'sem SKU')}): {str(e)}")

        db.commit()
        return {
            "success": True,
            "created_count": created_count,
            "updated_count": updated_count,
            "errors": errors,
        }

    @staticmethod
    def import_customers(db: Session, file_content: bytes, filename: str) -> Dict[str, Any]:
        """Lê arquivo Excel (.xlsx) ou CSV e realiza a importação de clientes."""
        rows = SystemService._read_spreadsheet(file_content, filename)
        if not rows:
            raise HTTPException(status_code=400, detail="Planilha vazia ou sem dados para importar.")

        created_count = 0
        updated_count = 0
        errors = []

        for idx, row in enumerate(rows, start=2):
            try:
                name = str(row.get("Nome Completo") or "").strip()
                if not name:
                    continue

                phone = str(row.get("Telefone / WhatsApp") or "").strip() or None
                email = str(row.get("E-mail") or "").strip().lower() or None
                nickname = str(row.get("Apelido") or "").strip() or None
                city = str(row.get("Cidade") or "").strip() or None
                state_uf = str(row.get("Estado") or "").strip().upper()[:2] or None
                address_line = str(row.get("Endereço") or "").strip() or None
                neighborhood = str(row.get("Bairro") or "").strip() or None
                zip_code = str(row.get("CEP") or "").strip() or None
                origin_channel = str(row.get("Canal de Origem") or "").strip() or None
                birth_date = str(row.get("Data de Nascimento") or "").strip() or None
                tags = str(row.get("Tags") or "").strip() or None

                # Buscar existente por telefone ou email
                existing = None
                if phone:
                    existing = db.query(Customer).filter(Customer.phone == phone).first()
                if not existing and email:
                    existing = db.query(Customer).filter(Customer.email == email).first()

                if existing:
                    existing.name = name
                    if nickname: existing.nickname = nickname
                    if city: existing.city = city
                    if state_uf: existing.state = state_uf
                    if address_line: existing.address_line = address_line
                    if neighborhood: existing.neighborhood = neighborhood
                    if zip_code: existing.zip_code = zip_code
                    if origin_channel: existing.origin_channel = origin_channel
                    if birth_date: existing.birth_date = birth_date
                    if tags: existing.tags = tags
                    updated_count += 1
                else:
                    new_cust = Customer(
                        name=name,
                        phone=phone,
                        email=email,
                        nickname=nickname,
                        city=city,
                        state=state_uf,
                        address_line=address_line,
                        neighborhood=neighborhood,
                        zip_code=zip_code,
                        origin_channel=origin_channel,
                        birth_date=birth_date,
                        tags=tags,
                    )
                    db.add(new_cust)
                    created_count += 1
            except Exception as e:
                errors.append(f"Linha {idx} ({row.get('Nome Completo', 'sem nome')}): {str(e)}")

        db.commit()
        return {
            "success": True,
            "created_count": created_count,
            "updated_count": updated_count,
            "errors": errors,
        }

    # --- Manutenção & Diagnóstico ---
    @staticmethod
    def clean_orphaned_images(db: Session) -> Dict[str, Any]:
        """Varre o diretório de fotos no disco e exclui arquivos órfãos sem vínculo no banco."""
        media_dir = getattr(settings, "MEDIA_DIR", "/app/media")
        products_media_dir = os.path.join(media_dir, "products")
        if not os.path.exists(products_media_dir):
            return {"removed_count": 0, "freed_bytes": 0, "message": "Nenhum diretório de fotos encontrado."}

        # Conjunto de caminhos registrados no banco
        registered_files = set(
            row[0] for row in db.query(ProductImage.file_path).all() if row[0]
        )

        removed_count = 0
        freed_bytes = 0

        for root, _, files in os.walk(products_media_dir):
            for f in files:
                full_path = os.path.join(root, f)
                rel_path = os.path.relpath(full_path, media_dir).replace("\\", "/")
                # Se não estiver no banco
                if rel_path not in registered_files and full_path not in registered_files:
                    try:
                        size = os.path.getsize(full_path)
                        os.remove(full_path)
                        removed_count += 1
                        freed_bytes += size
                    except Exception:
                        pass

        return {
            "success": True,
            "removed_count": removed_count,
            "freed_bytes": freed_bytes,
            "freed_formatted": SystemService._format_bytes(freed_bytes),
        }

    @staticmethod
    def optimize_database(db: Session) -> Dict[str, Any]:
        """Executa OPTIMIZE TABLE nas tabelas para desfragmentação no MySQL."""
        tables = [
            "products", "stock_movements", "sales", "sale_items", "customers",
            "production_orders", "financial_transactions", "categories", "product_images"
        ]
        results = []
        for tbl in tables:
            try:
                res = db.execute(text(f"OPTIMIZE TABLE `{tbl}`;")).fetchall()
                results.append({"table": tbl, "status": "Otimizada"})
            except Exception as e:
                results.append({"table": tbl, "status": f"Ignorada ({e})"})
        return {"success": True, "details": results}

    @staticmethod
    def reset_demo_data(db: Session, admin_password: str, current_user_id: int) -> Dict[str, Any]:
        """Zera dados operacionais de teste (vendas, estoque, ordens) preservando cadastros base."""
        from app.models.user import User
        from app.core.security import verify_password

        user = db.get(User, current_user_id)
        if not user or not verify_password(admin_password, user.password_hash):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Senha de administrador incorreta. Ação não autorizada.",
            )

        try:
            db.execute(text("SET FOREIGN_KEY_CHECKS = 0;"))
            # Limpar vendas, itens, ordens e movimentos de estoque
            db.execute(text("DELETE FROM sale_items;"))
            db.execute(text("DELETE FROM sales;"))
            db.execute(text("DELETE FROM stock_movements;"))
            db.execute(text("DELETE FROM production_order_items;"))
            db.execute(text("DELETE FROM production_orders;"))
            db.execute(text("DELETE FROM financial_transactions;"))
            db.execute(text("SET FOREIGN_KEY_CHECKS = 1;"))
            db.commit()
            return {"success": True, "message": "Dados de teste zerados com sucesso! Cadastros base foram preservados."}
        except Exception as e:
            db.rollback()
            db.execute(text("SET FOREIGN_KEY_CHECKS = 1;"))
            raise HTTPException(status_code=500, detail=f"Erro ao zerar dados de teste: {e}")

    @staticmethod
    def get_system_status(db: Session) -> Dict[str, Any]:
        """Retorna indicadores de status técnico, contagens de registros e armazenamento."""
        media_dir = getattr(settings, "MEDIA_DIR", "/app/media")
        total_media_size = 0
        media_files_count = 0
        if os.path.exists(media_dir):
            for root, _, files in os.walk(media_dir):
                for f in files:
                    try:
                        total_media_size += os.path.getsize(os.path.join(root, f))
                        media_files_count += 1
                    except Exception:
                        pass

        # Contagens de registros
        stats = {
            "products_total": db.query(Product).count(),
            "products_finished": db.query(Product).filter(Product.item_type == "PRODUTO_ACABADO", Product.parent_id == None).count(),
            "variations_count": db.query(Product).filter(Product.parent_id != None).count(),
            "customers_count": db.query(Customer).count(),
            "sales_count": db.query(Sale).count(),
            "stock_movements_count": db.query(StockMovement).count(),
            "recipes_count": db.query(ProductionRecipe).count(),
            "orders_count": db.query(ProductionOrder).count(),
            "financial_count": db.query(FinancialTransaction).count(),
        }

        # Status do Banco
        db_status = "Online / Saudável"
        try:
            db.execute(text("SELECT 1;"))
        except Exception:
            db_status = "Erro de conexão"

        return {
            "version": getattr(settings, "VERSION", "0.1.0"),
            "app_name": getattr(settings, "PROJECT_NAME", "Casa Luma Estoque"),
            "database_status": db_status,
            "media_storage_size_bytes": total_media_size,
            "media_storage_formatted": SystemService._format_bytes(total_media_size),
            "media_files_count": media_files_count,
            "counts": stats,
            "server_time": datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M:%S UTC"),
        }

    # --- Utilitários Internos ---
    @staticmethod
    def _create_spreadsheet(
        title: str,
        headers: List[str],
        rows: List[List[Any]],
        export_format: str,
        filename_prefix: str
    ) -> Tuple[io.BytesIO, str, str]:
        """Cria um arquivo .xlsx estilizado ou .csv com UTF-8 BOM."""
        now_str = datetime.now().strftime("%Y%m%d_%H%M")
        
        if export_format.lower() == "csv" or not HAS_OPENPYXL:
            buffer = io.BytesIO()
            # Adicionar UTF-8 BOM para o Excel abrir sem problemas de acentuação
            buffer.write(b"\xef\xbb\xbf")
            text_buffer = io.StringIO()
            writer = csv.writer(text_buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
            writer.writerow(headers)
            for r in rows:
                writer.writerow([str(v) if v is not None else "" for v in r])
            buffer.write(text_buffer.getvalue().encode("utf-8"))
            buffer.seek(0)
            filename = f"{filename_prefix}_{now_str}.csv"
            media_type = "text/csv; charset=utf-8"
            return buffer, filename, media_type
        else:
            # Excel .xlsx de alta qualidade
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = title[:30]

            # Estilos
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="6B4C35", end_color="6B4C35", fill_type="solid")
            border_thin = Border(
                left=Side(style='thin', color='E5E7EB'),
                right=Side(style='thin', color='E5E7EB'),
                top=Side(style='thin', color='E5E7EB'),
                bottom=Side(style='thin', color='E5E7EB')
            )

            # Cabeçalhos
            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Linhas de dados
            for row_idx, r_data in enumerate(rows, 2):
                fill_color = "FDFBF7" if row_idx % 2 == 0 else "FFFFFF"
                row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                for col_idx, val in enumerate(r_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.fill = row_fill
                    cell.border = border_thin
                    if isinstance(val, (int, float, Decimal)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Ajuste automático de largura de coluna
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

            ws.row_dimensions[1].height = 28

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            filename = f"{filename_prefix}_{now_str}.xlsx"
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return buffer, filename, media_type

    @staticmethod
    def _read_spreadsheet(file_content: bytes, filename: str) -> List[Dict[str, Any]]:
        """Lê arquivo .xlsx ou .csv e retorna lista de dicionários mapeando cabeçalhos."""
        rows: List[Dict[str, Any]] = []

        if filename.endswith(".xlsx"):
            if not HAS_OPENPYXL:
                raise HTTPException(status_code=500, detail="Módulo openpyxl não disponível para ler arquivo .xlsx.")
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
            headers = []
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=1, column=col).value
                headers.append(str(val).strip() if val is not None else f"col_{col}")

            for row_idx in range(2, ws.max_row + 1):
                row_dict = {}
                has_content = False
                for col_idx, h in enumerate(headers, 1):
                    c_val = ws.cell(row=row_idx, column=col_idx).value
                    if c_val is not None and str(c_val).strip() != "":
                        has_content = True
                    row_dict[h] = c_val
                if has_content:
                    rows.append(row_dict)
        else:
            # CSV (tenta detectar delimitador ';' ou ',')
            text_content = file_content.decode("utf-8-sig", errors="replace")
            sample = text_content[:2048]
            delimiter = ";" if ";" in sample else ","
            reader = csv.DictReader(io.StringIO(text_content), delimiter=delimiter)
            for r in reader:
                clean_r = {k.strip(): v.strip() if isinstance(v, str) else v for k, v in r.items() if k}
                if any(clean_r.values()):
                    rows.append(clean_r)

        return rows

    @staticmethod
    def _parse_decimal(val: Any) -> Optional[Decimal]:
        if val is None or str(val).strip() == "":
            return None
        s = str(val).strip().replace("R$", "").replace(" ", "")
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".")
        elif "," in s:
            s = s.replace(",", ".")
        try:
            return Decimal(s)
        except Exception:
            return None

    @staticmethod
    def _row_to_dict(obj: Any) -> Dict[str, Any]:
        """Converte uma instância SQLAlchemy em dicionário serializável."""
        d = {}
        for col in obj.__table__.columns:
            val = getattr(obj, col.name)
            if isinstance(val, (datetime, date)):
                d[col.name] = val.isoformat()
            elif isinstance(val, Decimal):
                d[col.name] = str(val)
            elif hasattr(val, "value"):
                d[col.name] = val.value
            else:
                d[col.name] = val
        return d

    @staticmethod
    def _format_bytes(bytes_val: int) -> str:
        for unit in ["B", "KB", "MB", "GB"]:
            if bytes_val < 1024.0:
                return f"{bytes_val:.1f} {unit}"
            bytes_val /= 1024.0
        return f"{bytes_val:.1f} TB"


system_service = SystemService()
